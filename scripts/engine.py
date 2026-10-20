#!/usr/bin/env python3
"""Listwright CLI.

Each step reads and writes files inside data/work/<run>/ so every step can be
re-run and inspected. Output is JSON on stdout for the slash commands to read;
problems are printed as a plain-English error plus one next step.

  load <file>                 copy the upload to data/raw/, suggest a column mapping
  save-pull <folder>          turn browser-pulled pages into a run (instead of load)
  clean     [--run R]         apply data/work/<run>/mapping.json, clean every lead
  dedupe    [--run R]         remove duplicates, suppressed and already-contacted leads
  score     [--run R]         score against config/icp_weights.json
  threshold [--run R] --min-score N   keep leads scoring N or more -> leads.csv
  summary   [--run R]         print the run's stats

  pending-enrich [--run R]    leads with no email/phone yet, with lookup fields
  merge-enrichment <file>     write email/phone/source returned by an enrichment tool
  pending-drafts [--run R] [--limit N]   next leads that still need emails
  add-drafts <file>           save drafted touches (JSON) for those leads
  export --name NAME [--run R]          write data/campaigns/<name>_<date>.csv
"""

import argparse
import csv
import datetime as dt
import glob
import io
import json
import os
import re
import shutil
import sys
import unicodedata
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import cleaning as cl  # noqa: E402
import scoring  # noqa: E402
from schema import BASE_COLUMNS, MAPPABLE_FIELDS, WORKING_COLUMNS, suggest_mapping  # noqa: E402
from schema import touch_columns as schema_touch_columns  # noqa: E402

ROOT = Path(os.environ.get("LISTWRIGHT_ROOT") or Path(__file__).resolve().parent.parent)
DATA = ROOT / "data"
RAW, WORK, CAMPAIGNS = DATA / "raw", DATA / "work", DATA / "campaigns"
SUPPRESSION = DATA / "suppression.csv"
ICP = ROOT / "config" / "icp_weights.json"
SEQUENCE = ROOT / "config" / "sequence.json"

LEAD_COLUMNS = [c for c in BASE_COLUMNS if c not in ("icp_score", "score_reasons")] + WORKING_COLUMNS
SCORED_COLUMNS = LEAD_COLUMNS + ["icp_score", "score_reasons"]
FINAL_COLUMNS = ["lead_id"] + SCORED_COLUMNS


class UserError(Exception):
    def __init__(self, message, next_step):
        super().__init__(message)
        self.next_step = next_step


# --- file io ----------------------------------------------------------------

def clean_path(p):
    """Handle drag-and-drop paths: surrounding quotes, backslash-escaped spaces, ~."""
    p = p.strip().strip("'\"")
    p = re.sub(r"\\(.)", r"\1", p)
    return Path(os.path.expanduser(p))


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (dt.datetime, dt.date)):
        return v.date().isoformat() if isinstance(v, dt.datetime) else v.isoformat()
    return str(v).strip()


def _trim_table(rows):
    rows = [[_cell(v) for v in r] for r in rows]
    # Skip preamble lines above the real header (needs 2+ filled cells).
    start = next((i for i, r in enumerate(rows) if sum(1 for v in r if v) >= 2), None)
    if start is None:
        return [], []
    headers = rows[start]
    while headers and not headers[-1]:
        headers = headers[:-1]
    seen, uniq = {}, []
    for i, h in enumerate(headers):
        h = h or f"Column {i + 1}"
        seen[h] = seen.get(h, 0) + 1
        uniq.append(h if seen[h] == 1 else f"{h} ({seen[h]})")
    body = []
    for r in rows[start + 1:]:
        r = (r + [""] * len(uniq))[:len(uniq)]
        if any(r):
            body.append(dict(zip(uniq, r)))
    return uniq, body


def read_table(path):
    ext = path.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        try:
            import openpyxl
        except ImportError:
            raise UserError("Excel support isn't installed.", "Run ./install.sh, then try again.")
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except Exception:
            raise UserError("That Excel file couldn't be opened. It may be damaged or password-protected.",
                            "Open it in Excel, save a fresh copy as .xlsx or .csv, and try again.")
        sheets = []
        for ws in wb.worksheets:
            headers, body = _trim_table(ws.iter_rows(values_only=True))
            if body:
                sheets.append((ws.title, headers, body))
        wb.close()
        if not sheets:
            return [], [], []
        notes = []
        if len(sheets) > 1:
            notes.append(f"File has {len(sheets)} sheets with data; using the first one, '{sheets[0][0]}'.")
        return sheets[0][1], sheets[0][2], notes
    if ext == ".xls":
        raise UserError("Old-style .xls files aren't supported.",
                        "Open it in Excel and save it as .xlsx or .csv, then try again.")
    if ext not in (".csv", ".tsv", ".txt"):
        raise UserError(f"'{path.name}' isn't a spreadsheet this tool can read.",
                        "Use a .csv or .xlsx file.")
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            content = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = content[:8192]
    try:
        delim = csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
    except csv.Error:
        delim = "\t" if ext == ".tsv" else ","
    headers, body = _trim_table(csv.reader(io.StringIO(content), delimiter=delim))
    return headers, body, []


def write_csv(path, columns, rows, excel=False):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8-sig" if excel else "utf-8") as f:
        w = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow({c: r.get(c, "") for c in columns})


def read_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


# --- runs -------------------------------------------------------------------

def run_dir(run=None):
    if run:
        d = WORK / run
        if not d.is_dir():
            raise UserError(f"There's no run called '{run}'.", "Run /ingest-rank to start a new one.")
        return d
    runs = sorted(p for p in WORK.glob("*") if p.is_dir()) if WORK.exists() else []
    if not runs:
        raise UserError("No leads have been loaded yet.", "Run /ingest-rank and upload a file.")
    return runs[-1]


def need(d, name, step):
    p = d / name
    if not p.exists():
        raise UserError(f"This step needs {name}, which doesn't exist yet.", f"Run the '{step}' step first.")
    return p


def load_stats(d):
    p = d / "stats.json"
    return json.loads(p.read_text()) if p.exists() else {}


def save_stats(d, **kw):
    stats = load_stats(d)
    stats.update(kw)
    (d / "stats.json").write_text(json.dumps(stats, indent=2))
    return stats


# --- commands ---------------------------------------------------------------

def new_run_dir():
    run = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    d = WORK / run
    n = 2
    while d.exists():
        d = WORK / f"{run}-{n}"
        n += 1
    d.mkdir(parents=True)
    return d


PULL_FIELDS = ["full_name", "title", "company", "location", "sales_nav_url", "linkedin_url",
               "tenure", "recent_job_change", "headline", "industry", "company_size", "seniority"]


def cmd_save_pull(args):
    """Combine browser-pulled pages (page-*.json in a pull folder) into a new run.

    Each page file is a JSON list of leads using PULL_FIELDS names. The combined
    CSV is saved to data/raw/ as the untouched input, and the mapping is
    written automatically because the columns are already standard.
    """
    folder = clean_path(args.folder)
    pages = sorted(folder.glob("page-*.json"), key=lambda p: int(re.sub(r"\D", "", p.stem) or 0))
    if not pages:
        raise UserError("No pulled pages were found.", "Pull at least one page with /ingest-rank, then try again.")
    rows, seen, skipped = [], set(), 0
    for page in pages:
        try:
            items = json.loads(page.read_text())
        except json.JSONDecodeError:
            raise UserError(f"{page.name} is damaged.", "Delete that page file and pull it again.")
        for item in items:
            row = {f: str(item.get(f) or "").strip() for f in PULL_FIELDS}
            key = cl.url_key(row["sales_nav_url"]) or (row["full_name"].lower(), row["company"].lower())
            if not row["full_name"] or key in seen:
                skipped += 1
                continue
            seen.add(key)
            rows.append(row)
    if not rows:
        raise UserError("The pulled pages had no leads in them.", "Check the list URL and try again.")

    d = new_run_dir()
    RAW.mkdir(parents=True, exist_ok=True)
    raw_copy = RAW / f"{d.name}_salesnav_pull.csv"
    write_csv(raw_copy, PULL_FIELDS, rows)
    write_csv(d / "loaded.csv", PULL_FIELDS, rows)
    mapping = {f: (f if f in PULL_FIELDS else None) for f in MAPPABLE_FIELDS}
    (d / "mapping.suggested.json").write_text(json.dumps(mapping, indent=2))
    (d / "mapping.json").write_text(json.dumps(mapping, indent=2))
    save_stats(d, run=d.name, source_file=f"Sales Navigator ({len(pages)} pages)",
               raw_copy=str(raw_copy), total_in=len(rows))
    out(run=d.name, raw_copy=str(raw_copy.relative_to(ROOT)), pages=len(pages), rows=len(rows),
        skipped_repeats_or_blank=skipped, sample=rows[:3])


def cmd_load(args):
    src = clean_path(args.file)
    if not src.exists():
        raise UserError(f"Couldn't find the file '{src}'.",
                        "Drag the file into this window again, or paste its full path.")
    if src.is_dir():
        raise UserError("That's a folder, not a file.", "Drag in the .csv or .xlsx file itself.")
    headers, rows, notes = read_table(src)
    if not rows:
        raise UserError("That file has no leads in it.", "Check you exported the right list and try again.")

    d = new_run_dir()
    run = d.name
    RAW.mkdir(parents=True, exist_ok=True)
    raw_copy = RAW / f"{run}_{src.name}"
    shutil.copy2(src, raw_copy)  # untouched original

    write_csv(d / "loaded.csv", headers, rows)
    mapping, unmapped = suggest_mapping(headers)
    (d / "mapping.suggested.json").write_text(json.dumps(mapping, indent=2))

    warnings = list(notes)
    if not (mapping["full_name"] or mapping["first_name"]):
        warnings.append("No name column found. Pick which column has the person's name.")
    if not (mapping["linkedin_url"] or mapping["sales_nav_url"]):
        warnings.append("No LinkedIn URL column found. Dedupe will rely on name + company only.")
    if not mapping["company"]:
        warnings.append("No company column found.")
    contact_cols = [h for h in headers if re.search(r"e-?mail|phone|mobile", h, re.I)]
    if contact_cols:
        warnings.append("Email/phone columns in the file are ignored. Contact info only comes "
                        "from /enrich: " + ", ".join(contact_cols))

    save_stats(d, run=run, source_file=str(src), raw_copy=str(raw_copy), total_in=len(rows))
    out(run=run, raw_copy=str(raw_copy.relative_to(ROOT)), rows=len(rows), headers=headers,
        sample=rows[:3], suggested_mapping=mapping, unmapped_columns=unmapped,
        mapping_file=str((d / "mapping.json").relative_to(ROOT)), warnings=warnings)


def _pick(row, mapping, field):
    h = mapping.get(field)
    return row.get(h, "") if h else ""


def clean_row(row, mapping):
    first = cl.clean_name_part(_pick(row, mapping, "first_name"))
    last = cl.clean_name_part(_pick(row, mapping, "last_name"))
    if not (first or last):
        first, last = cl.split_name(cl.clean_name_part(_pick(row, mapping, "full_name")))
    elif first and not last and " " in first:
        first, last = cl.split_name(first)

    location, city, region, country = cl.parse_location(_pick(row, mapping, "location"))

    li = cl.normalize_url(_pick(row, mapping, "linkedin_url"))
    sn = cl.normalize_url(_pick(row, mapping, "sales_nav_url"))
    if li and cl.is_sales_nav_url(li):
        sn, li = sn or li, ""
    if sn and not cl.is_sales_nav_url(sn) and "/in/" in sn and not li:
        li, sn = sn, ""

    tenure = cl.text(_pick(row, mapping, "tenure"))
    return {
        "first_name": first,
        "last_name": last,
        "full_name": " ".join(p for p in (first, last) if p),
        "title": cl.normalize_title(_pick(row, mapping, "title")),
        "company": cl.normalize_company(_pick(row, mapping, "company")),
        "location": location,
        "linkedin_url": li,
        "sales_nav_url": sn,
        "tenure": tenure,
        "recent_job_change": cl.yes_no(_pick(row, mapping, "recent_job_change")),
        "email": "", "phone": "", "enrichment_source": "",
        "headline": cl.text(_pick(row, mapping, "headline")),
        "industry": cl.text(_pick(row, mapping, "industry")),
        "company_size": cl.text(_pick(row, mapping, "company_size")),
        "seniority": cl.text(_pick(row, mapping, "seniority")),
        "city": city, "region": region, "country": country,
        "tenure_months": cl.tenure_months(tenure),
    }


def cmd_clean(args):
    d = run_dir(args.run)
    loaded = read_csv(need(d, "loaded.csv", "load"))
    mp = d / "mapping.json"
    if not mp.exists():
        raise UserError("The column mapping hasn't been confirmed yet.",
                        "Confirm the mapping so it's saved as mapping.json, then clean again.")
    mapping = json.loads(mp.read_text())
    unknown = [f for f in mapping if f not in MAPPABLE_FIELDS]
    if unknown:
        raise UserError(f"The mapping has fields this tool doesn't know: {', '.join(unknown)}.",
                        f"Use only: {', '.join(MAPPABLE_FIELDS)}.")
    headers = set(loaded[0].keys()) if loaded else set()
    missing = [h for h in mapping.values() if h and h not in headers]
    if missing:
        raise UserError(f"The mapping points to columns that aren't in the file: {', '.join(missing)}.",
                        "Fix the mapping and try again.")
    if not (mapping.get("full_name") or mapping.get("first_name")):
        raise UserError("No column is mapped to the person's name.", "Map a name column and try again.")

    cleaned, before_after, empty = [], [], 0
    for row in loaded:
        lead = clean_row(row, mapping)
        if not (lead["full_name"] or lead["linkedin_url"] or lead["sales_nav_url"]):
            empty += 1
            continue
        cleaned.append(lead)
        raw_name = _pick(row, mapping, "full_name") or " ".join(
            filter(None, (_pick(row, mapping, "first_name"), _pick(row, mapping, "last_name"))))
        changed = {
            "name": (raw_name, lead["full_name"]),
            "title": (_pick(row, mapping, "title"), lead["title"]),
            "company": (_pick(row, mapping, "company"), lead["company"]),
        }
        changed = {k: {"before": a, "after": b} for k, (a, b) in changed.items() if a.strip() != b}
        if changed and len(before_after) < 5:
            before_after.append(changed)

    write_csv(d / "cleaned.csv", LEAD_COLUMNS, cleaned)
    save_stats(d, cleaned=len(cleaned), empty_rows_dropped=empty)
    out(run=d.name, cleaned=len(cleaned), empty_rows_dropped=empty, examples=before_after)


def _fold(s):
    s = unicodedata.normalize("NFKD", s or "")
    return re.sub(r"[^a-z]", "", "".join(c for c in s if not unicodedata.combining(c)).lower())


def lead_keys(r):
    """Identity keys: each URL, plus name+company when both exist."""
    keys = []
    for f in ("linkedin_url", "sales_nav_url"):
        k = cl.url_key(r.get(f, ""))
        if k:
            keys.append("url:" + k)
    name = _fold(r.get("full_name") or f"{r.get('first_name', '')}{r.get('last_name', '')}")
    comp = cl.company_key(r.get("company", ""))
    if name and comp:
        keys.append(f"nc:{name}|{comp}")
    return keys


def _key_index(rows, label):
    idx = {}
    for r in rows:
        for k in lead_keys(r):
            idx.setdefault(k, label)
    return idx


def cmd_dedupe(args):
    d = run_dir(args.run)
    leads = read_csv(need(d, "cleaned.csv", "clean"))

    suppressed_idx = _key_index(read_csv(SUPPRESSION), "suppression list") if SUPPRESSION.exists() else {}
    past_idx = {}
    for f in sorted(glob.glob(str(CAMPAIGNS / "*.csv"))):
        past_idx.update({k: v for k, v in _key_index(read_csv(Path(f)), Path(f).name).items()
                         if k not in past_idx})

    kept, removed = [], []
    seen = {}
    counts = {"duplicate_url": 0, "duplicate_name_company": 0, "suppressed": 0, "already_in_campaign": 0}
    for r in leads:
        keys = lead_keys(r)
        hit = next((k for k in keys if k in suppressed_idx), None)
        if hit:
            counts["suppressed"] += 1
            removed.append({**r, "removed_reason": "suppressed", "matched_on": hit})
            continue
        hit = next((k for k in keys if k in past_idx), None)
        if hit:
            counts["already_in_campaign"] += 1
            removed.append({**r, "removed_reason": f"already in {past_idx[hit]}", "matched_on": hit})
            continue
        hit = next((k for k in keys if k in seen), None)
        if hit:
            keeper = kept[seen[hit]]
            for c, v in r.items():  # fill gaps in the kept copy, never overwrite
                if v and not keeper.get(c):
                    keeper[c] = v
            for k in keys:
                seen.setdefault(k, seen[hit])
            counts["duplicate_url" if hit.startswith("url:") else "duplicate_name_company"] += 1
            removed.append({**r, "removed_reason": "duplicate", "matched_on": hit})
            continue
        for k in keys:
            seen[k] = len(kept)
        kept.append(r)

    write_csv(d / "deduped.csv", LEAD_COLUMNS, kept)
    write_csv(d / "removed.csv", LEAD_COLUMNS + ["removed_reason", "matched_on"], removed)
    save_stats(d, after_dedupe=len(kept), **counts)
    out(run=d.name, kept=len(kept), **counts)


def load_icp():
    if not ICP.exists():
        raise UserError("There's no ICP scoring set up yet.", "Run /configure-icp first.")
    try:
        cfg = json.loads(ICP.read_text())
        scoring.validate(cfg)
    except json.JSONDecodeError:
        raise UserError("The ICP file is damaged.", "Run /configure-icp to set it up again.")
    except scoring.ConfigError as e:
        raise UserError(str(e), "Run /configure-icp to fix it.")
    return cfg


BUCKETS = [(80, 100), (60, 79), (40, 59), (20, 39), (0, 19)]


def histogram(scores):
    return {f"{lo}-{hi}": sum(1 for s in scores if lo <= s <= hi) for lo, hi in BUCKETS}


def cmd_score(args):
    d = run_dir(args.run)
    leads = read_csv(need(d, "deduped.csv", "dedupe"))
    cfg = load_icp()
    for r in leads:
        r["icp_score"], r["score_reasons"] = scoring.score_lead(r, cfg)
    leads.sort(key=lambda r: -r["icp_score"])
    write_csv(d / "scored.csv", SCORED_COLUMNS, leads)
    scores = [r["icp_score"] for r in leads]
    stats = save_stats(d, scored=len(leads), score_distribution=histogram(scores),
                       average_score=round(sum(scores) / len(scores), 1) if scores else 0)
    top = [{k: r[k] for k in ("full_name", "title", "company", "icp_score", "score_reasons")}
           for r in leads[:5]]
    out(run=d.name, scored=len(leads), score_distribution=stats["score_distribution"],
        average_score=stats["average_score"], top_leads=top)


def cmd_threshold(args):
    d = run_dir(args.run)
    leads = read_csv(need(d, "scored.csv", "score"))
    existing = d / "leads.csv"
    if (d / "drafts.json").exists() or (
            existing.exists() and any(r["email"] or r["phone"] for r in read_csv(existing))):
        raise UserError("This run already has drafted emails or enriched contacts, and changing the "
                        "cutoff now would throw that work away.",
                        "Keep going with this list, or run /ingest-rank again to start a fresh run.")
    kept = [r for r in leads if int(r["icp_score"] or 0) >= args.min_score]
    for i, r in enumerate(kept, 1):
        r["lead_id"] = f"L{i:04d}"
    write_csv(d / "leads.csv", FINAL_COLUMNS, kept)
    save_stats(d, min_score=args.min_score, final_leads=len(kept), below_threshold=len(leads) - len(kept))
    out(run=d.name, min_score=args.min_score, kept=len(kept), dropped=len(leads) - len(kept),
        leads_file=str((d / "leads.csv").relative_to(ROOT)))


def cmd_summary(args):
    d = run_dir(args.run)
    out(**load_stats(d))


# --- enrichment -------------------------------------------------------------

EMAIL_RE = re.compile(r"^[^@\s,;<>]+@[^@\s,;<>]+\.[a-z]{2,}$", re.I)


def read_json_input(path):
    p = clean_path(path)
    if not p.exists():
        raise UserError(f"Couldn't find '{p}'.", "Write the JSON file first, then run this again.")
    try:
        return json.loads(p.read_text())
    except json.JSONDecodeError as e:
        raise UserError(f"'{p.name}' isn't valid JSON ({e.msg}, line {e.lineno}).", "Fix the file and try again.")


def cmd_pending_enrich(args):
    d = run_dir(args.run)
    leads = read_csv(need(d, "leads.csv", "threshold"))
    fields = ("lead_id", "first_name", "last_name", "title", "company", "linkedin_url", "sales_nav_url")
    todo = [{f: r[f] for f in fields} for r in leads if not (r["email"] or r["phone"])]
    out(run=d.name, total_leads=len(leads), need_lookup=len(todo), leads=todo)


def cmd_merge_enrichment(args):
    """Fill email/phone only from tool results; never overwrite, never guess (hard rule 2)."""
    d = run_dir(args.run)
    path = need(d, "leads.csv", "threshold")
    leads = read_csv(path)
    by_id = {r["lead_id"]: r for r in leads}
    data = read_json_input(args.file)
    default_source = str(data.get("source", "")).strip() if isinstance(data, dict) else ""
    results = data.get("results", []) if isinstance(data, dict) else data

    added = {"email": 0, "phone": 0}
    rejected = []
    for res in results:
        lead = by_id.get(str(res.get("lead_id", "")))
        if lead is None:
            rejected.append({"lead_id": res.get("lead_id"), "why": "unknown lead_id"})
            continue
        source = str(res.get("source") or default_source).strip()
        if not source:
            rejected.append({"lead_id": lead["lead_id"], "why": "no enrichment source named"})
            continue
        email = str(res.get("email") or "").strip()
        phone = str(res.get("phone") or "").strip()
        wrote = False
        if email:
            if not EMAIL_RE.match(email):
                rejected.append({"lead_id": lead["lead_id"], "why": f"not a valid email: {email}"})
            elif not lead["email"]:
                lead["email"], wrote = email, True
                added["email"] += 1
        if phone:
            if not 7 <= len(re.sub(r"\D", "", phone)) <= 15:
                rejected.append({"lead_id": lead["lead_id"], "why": f"not a valid phone: {phone}"})
            elif not lead["phone"]:
                lead["phone"], wrote = phone, True
                added["phone"] += 1
        if wrote:
            sources = [s for s in lead["enrichment_source"].split("; ") if s]
            if source not in sources:
                lead["enrichment_source"] = "; ".join(sources + [source])

    write_csv(path, FINAL_COLUMNS, leads)
    with_email = sum(1 for r in leads if r["email"])
    with_phone = sum(1 for r in leads if r["phone"])
    save_stats(d, enriched_emails=with_email, enriched_phones=with_phone)
    out(run=d.name, emails_added=added["email"], phones_added=added["phone"],
        leads_with_email=with_email, leads_with_phone=with_phone, total_leads=len(leads),
        rejected=rejected)


# --- drafting & export ------------------------------------------------------

def load_sequence():
    if not SEQUENCE.exists():
        raise UserError("Your email sequence isn't set up yet.", "Run /setup-sequence first.")
    try:
        cfg = json.loads(SEQUENCE.read_text())
    except json.JSONDecodeError:
        raise UserError("The sequence file is damaged.", "Run /setup-sequence to set it up again.")
    touches = cfg.get("touches")
    if not isinstance(touches, list) or not 1 <= len(touches) <= 10:
        raise UserError("The sequence needs between 1 and 10 touches.", "Run /setup-sequence to fix it.")
    days = [t.get("day") for t in touches]
    if not all(isinstance(x, int) and x >= 1 for x in days) or days != sorted(set(days)):
        raise UserError("Touch days must be whole numbers that go up, starting at day 1 or later.",
                        "Run /setup-sequence to fix it.")
    return cfg


PLACEHOLDER_RE = re.compile(r"\{\{.*?\}\}|\[(?:first ?name|last ?name|name|company|title|your name)\]|"
                            r"<(?:first_?name|company|name)>|\bINSERT\b|\bTODO\b|\bXX+\b", re.I)


def _drafts_file(d):
    return d / "drafts.json"


def load_drafts(d):
    p = _drafts_file(d)
    return json.loads(p.read_text()) if p.exists() else {}


DRAFT_FIELDS = ("lead_id", "first_name", "last_name", "title", "company", "location", "city",
                "region", "country", "tenure", "recent_job_change", "headline", "industry",
                "company_size", "seniority")


def cmd_pending_drafts(args):
    d = run_dir(args.run)
    leads = read_csv(need(d, "leads.csv", "threshold"))
    drafts = load_drafts(d)
    todo = [r for r in leads if r["lead_id"] not in drafts]
    if args.ids:
        wanted = set(args.ids.split(","))
        todo = [r for r in leads if r["lead_id"] in wanted]
    batch = [{f: r[f] for f in DRAFT_FIELDS if r.get(f)} for r in todo[:args.limit]]
    out(run=d.name, total_leads=len(leads), drafted=len(drafts), remaining=len(todo), leads=batch)


def cmd_add_drafts(args):
    d = run_dir(args.run)
    leads = {r["lead_id"]: r for r in read_csv(need(d, "leads.csv", "threshold"))}
    n = len(load_sequence()["touches"])
    data = read_json_input(args.file)
    items = data.get("drafts", []) if isinstance(data, dict) else data
    drafts = load_drafts(d)

    saved, rejected = [], []
    for item in items:
        lid = str(item.get("lead_id", ""))
        touches = item.get("touches") or []
        problems = []
        if lid not in leads:
            problems.append("unknown lead_id")
        if len(touches) != n:
            problems.append(f"has {len(touches)} touches, sequence needs {n}")
        for i, t in enumerate(touches, 1):
            subject, body = str(t.get("subject") or "").strip(), str(t.get("body") or "").strip()
            if not body:
                problems.append(f"touch {i} has no body")
            if i == 1 and not subject:
                problems.append("touch 1 has no subject")
            left = PLACEHOLDER_RE.findall(subject + "\n" + body)
            if left:
                problems.append(f"touch {i} has unfilled placeholders: {', '.join(sorted(set(left)))}")
        if problems:
            rejected.append({"lead_id": lid, "problems": problems})
            continue
        drafts[lid] = [{"subject": str(t.get("subject") or "").strip(),
                        "body": str(t["body"]).strip()} for t in touches]
        saved.append(lid)

    _drafts_file(d).write_text(json.dumps(drafts, indent=1, ensure_ascii=False))
    out(run=d.name, saved=len(saved), rejected=rejected, drafted=len(drafts),
        remaining=sum(1 for lid in leads if lid not in drafts))


def _safe_cell(v):
    """Stop spreadsheet apps from running cell text as a formula."""
    v = "" if v is None else str(v)
    if v[:1] in ("=", "+", "@", "\t", "\r") or (v[:1] == "-" and v[1:2] not in ("", " ")):
        return "'" + v
    return v


def _slug(name):
    return re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")[:60]


def cmd_export(args):
    d = run_dir(args.run)
    leads = read_csv(need(d, "leads.csv", "threshold"))
    seq = load_sequence()
    n = len(seq["touches"])
    drafts = load_drafts(d)
    missing = [r["lead_id"] for r in leads if r["lead_id"] not in drafts]
    if missing:
        raise UserError(f"{len(missing)} of {len(leads)} leads don't have emails yet.",
                        "Run /draft-sequence to finish them, then export again.")
    bad = [lid for lid, t in drafts.items() if len(t) != n and lid in {r["lead_id"] for r in leads}]
    if bad:
        raise UserError(f"{len(bad)} leads were drafted for a different number of touches than "
                        f"your sequence ({n}).", "Redraft those leads with /draft-sequence.")
    slug = _slug(args.name)
    if not slug:
        raise UserError("The campaign name needs at least one letter or number.", "Pick a short name, e.g. q4-cro-outreach.")

    columns = BASE_COLUMNS + schema_touch_columns(n)
    rows = []
    for r in leads:
        row = {c: r.get(c, "") for c in BASE_COLUMNS}
        for i, (cfg_t, t) in enumerate(zip(seq["touches"], drafts[r["lead_id"]]), 1):
            row[f"touch{i}_day"] = cfg_t["day"]
            row[f"touch{i}_subject"] = t["subject"]
            row[f"touch{i}_body"] = t["body"]
        rows.append({c: _safe_cell(v) for c, v in row.items()})

    CAMPAIGNS.mkdir(parents=True, exist_ok=True)
    base = f"{slug}_{dt.date.today().isoformat()}"
    k = 1
    while True:
        path = CAMPAIGNS / (f"{base}.csv" if k == 1 else f"{base}_{k}.csv")
        try:
            f = open(path, "x", newline="", encoding="utf-8-sig")  # 'x' never overwrites (hard rule 5)
            break
        except FileExistsError:
            k += 1
    with f:
        w = csv.DictWriter(f, fieldnames=columns)
        w.writeheader()
        w.writerows(rows)

    save_stats(d, exported_file=str(path.relative_to(ROOT)), exported_leads=len(rows))
    out(run=d.name, file=str(path.relative_to(ROOT)), full_path=str(path), leads=len(rows),
        touches=n, with_email=sum(1 for r in leads if r["email"]))


# --- main -------------------------------------------------------------------

def out(**kw):
    print(json.dumps(kw, indent=2, ensure_ascii=False, default=str))


def main(argv=None):
    p = argparse.ArgumentParser(prog="engine.py", description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("load")
    s.add_argument("file")
    for name in ("clean", "dedupe", "score", "summary"):
        sub.add_parser(name).add_argument("--run")
    s = sub.add_parser("threshold")
    s.add_argument("--run")
    s.add_argument("--min-score", type=int, required=True)
    sub.add_parser("pending-enrich").add_argument("--run")
    sub.add_parser("save-pull").add_argument("folder")
    for name in ("merge-enrichment", "add-drafts"):
        s = sub.add_parser(name)
        s.add_argument("file")
        s.add_argument("--run")
    s = sub.add_parser("pending-drafts")
    s.add_argument("--run")
    s.add_argument("--limit", type=int, default=20)
    s.add_argument("--ids", help="comma-separated lead_ids to show instead of the next batch")
    s = sub.add_parser("export")
    s.add_argument("--run")
    s.add_argument("--name", required=True)
    args = p.parse_args(argv)

    try:
        {"load": cmd_load, "clean": cmd_clean, "dedupe": cmd_dedupe, "score": cmd_score,
         "threshold": cmd_threshold, "summary": cmd_summary,
         "pending-enrich": cmd_pending_enrich, "merge-enrichment": cmd_merge_enrichment,
         "pending-drafts": cmd_pending_drafts, "add-drafts": cmd_add_drafts,
         "export": cmd_export, "save-pull": cmd_save_pull}[args.cmd](args)
    except UserError as e:
        print(f"ERROR: {e}\nNEXT STEP: {e.next_step}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
